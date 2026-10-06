import re
from difflib import SequenceMatcher

from docx import Document

from .base import Doc

_SKIP_STYLES = ("Heading", "Title", "Subtitle", "Caption", "TOC")
_HEADING = re.compile(r"(#{1,6})\s+(.*)")


def _paragraphs(document):
    seen, out = set(), []

    def add(paras):
        for p in paras:
            if p._p not in seen:
                seen.add(p._p)
                out.append(p)

    add(document.paragraphs)
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                add(cell.paragraphs)
    return out


def _locked(p):
    name = p.style.name if p.style is not None else ""
    return not p.text.strip() or name.startswith(_SKIP_STYLES) or bool(p._p.xpath("./w:hyperlink"))


def _merge(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def _patch(p, new):
    runs = p.runs
    olds = [r.text for r in runs]
    old = "".join(olds)
    if not runs or old != p.text:
        _merge(p, new)
        return
    bounds, pos = [], 0
    for t in olds:
        bounds.append((pos, pos + len(t)))
        pos += len(t)
    out = [""] * len(runs)
    for tag, i1, i2, j1, j2 in SequenceMatcher(None, old, new, autojunk=False).get_opcodes():
        if tag == "equal":
            for k, (s, e) in enumerate(bounds):
                lo, hi = max(s, i1), min(e, i2)
                if lo < hi:
                    out[k] += old[lo:hi]
        elif j2 > j1:
            k = next((k for k, (s, e) in enumerate(bounds) if s <= i1 < e), len(runs) - 1)
            out[k] += new[j1:j2]
    for r, t in zip(runs, out):
        if r.text != t:
            r.text = t


def read(path):
    paras = _paragraphs(Document(str(path)))
    return Doc(
        [p.text for p in paras],
        "docx",
        {i for i, p in enumerate(paras) if _locked(p)},
    )


def write(doc, units, out):
    document = Document(doc.meta["src"])
    for p, new in zip(_paragraphs(document), units):
        if new != p.text:
            _patch(p, new)
    document.save(str(out))


def write_new(doc, units, out):
    document = Document()
    for u in units:
        if not u.strip():
            continue
        m = _HEADING.fullmatch(u)
        if m:
            document.add_heading(m.group(2), len(m.group(1)))
        else:
            document.add_paragraph(u)
    document.save(str(out))
